# -*- coding: utf-8 -*-
"""
fingertap_pajarito.py
=====================
Minijuego de rehabilitación "Finger Tap" (pinza índice-pulgar) con un pájaro
flotante. Versión C: el pájaro sube si el paciente hace taps y cae si para.

Qué evalúa
----------
Simula la tarea de "finger tapping" de la escala MDS-UPDRS (bradicinesia):
el paciente junta y separa el índice y el pulgar el mayor número de veces
posible durante GAME_DURATION segundos. A partir de la cámara se obtienen:

    - Amplitud de cada tap (distancia máxima índice-pulgar entre taps, en cm)
    - Velocidad de cierre de cada tap (cm/s)
    - Intervalo entre taps (ITI, s)
    - Tendencia (pendiente de regresión lineal) de amplitud y velocidad:
      una pendiente negativa indica decremento (efecto secuencia)
    - Variabilidad de la amplitud (coeficiente de variación, %)
    - Temblor (frecuencia dominante y amplitud por eje) con la micro:bit,
      a través de ``temblor_sensor.py``

Flujo del programa
------------------
    1. Ventana Tkinter para registrar paciente (ID, género, edad).
    2. Bucle principal (cámara + MediaPipe Hands + dibujo con OpenCV):
         - Lee un frame, detecta la mano y calcula la distancia índice-pulgar.
         - Detecta cada tap (distancia < 1 cm) y actualiza la física del pájaro.
         - Dibuja el juego, el panel lateral con gráficas en tiempo real y la
           cámara en vivo.
    3. Al acabar el tiempo se detiene el sensor de temblor y se muestran los
       resultados en pantalla.
    4. Al pulsar Q se guardan: gráfica PNG, fila en Excel y gráfica FFT.

Nota: el script se ejecuta de arriba abajo (no hay función main()), por lo que
el estado del juego se guarda en variables globales.

Requiere: pip install opencv-python mediapipe==0.10.9 numpy matplotlib openpyxl
          pyserial scipy   (y ``temblor_sensor.py`` en la misma carpeta)
"""

import cv2
import mediapipe as mp
import time
import math
import numpy as np
import matplotlib.pyplot as plt
from datetime import datetime
import tkinter as tk
from tkinter import ttk, messagebox
from openpyxl import Workbook, load_workbook
from collections import deque
from temblor_sensor import TemblorSensor

# ================= TEMBLOR (micro:bit) =================
PUERTO_MICROBIT = "COM3"  # Cambiar por tu puerto
# El objeto se crea aquí pero la captura NO empieza hasta temblor.iniciar(),
# que se llama justo antes de arrancar el cronómetro de la partida.
temblor = TemblorSensor(puerto=PUERTO_MICROBIT)
GUARDAR_CSV_TEMBLOR = False  # CSV crudo (x,y,z por muestra). La frecuencia y
                              # amplitud por eje ya se guardan en el Excel de
                              # la sesion, asi que por defecto no se genera.

# ================= MEDIAPIPE =================
# MediaPipe Hands devuelve 21 landmarks (puntos) por mano, con coordenadas
# normalizadas (0-1) respecto al tamaño de la imagen.
mp_hands       = mp.solutions.hands
mp_draw        = mp.solutions.drawing_utils
mp_draw_styles = mp.solutions.drawing_styles

hands = mp_hands.Hands(
    max_num_hands=1,               # solo se evalúa una mano
    model_complexity=1,            # 0 = rápido, 1 = más preciso
    min_detection_confidence=0.7,  # confianza mínima para detectar la mano
    min_tracking_confidence=0.7    # confianza mínima para seguirla entre frames
)

# ================= VENTANA =================
# La ventana se divide en: zona de juego (izquierda) + panel de métricas (derecha).
WINDOW_W = 1280
WINDOW_H = 720
PANEL_W  = 400
GAME_W   = WINDOW_W - PANEL_W
WIN_NAME = "Rehabilitacion"
cv2.namedWindow(WIN_NAME, cv2.WINDOW_NORMAL)
cv2.resizeWindow(WIN_NAME, WINDOW_W, WINDOW_H)

# ================= CAMARA =================
cap = cv2.VideoCapture(0)   # 0 = cámara por defecto
cap.set(cv2.CAP_PROP_FRAME_WIDTH,  640)
cap.set(cv2.CAP_PROP_FRAME_HEIGHT, 480)

# ================= CONSTANTE ANATOMICA =================
# Referencia para convertir píxeles a centímetros: la distancia entre la muñeca
# (landmark 0) y la base del dedo medio (landmark 9) mide ~9,5 cm en un adulto.
# Esta escala se recalcula en cada frame, así el resultado no depende de lo
# cerca o lejos que esté la mano de la cámara.
HAND_REF_REAL_CM = 9.5   # muneca(0) -> base dedo medio(9): ~9.5 cm


# ================= PACIENTE =================
def obtener_datos_paciente():
    """Muestra una ventana Tkinter para registrar al paciente.

    Bloquea hasta que se pulsa "Iniciar Evaluacion" con datos válidos.
    Devuelve un diccionario ``{"id", "genero", "edad"}``. Si se cierra la
    ventana sin iniciar, devuelve un diccionario vacío (y el programa fallará
    al acceder a ``paciente["id"]``).
    """
    datos = {}

    ventana = tk.Tk()
    ventana.title("Registro de Paciente")
    ventana.geometry("420x300")
    ventana.resizable(False, False)

    tk.Label(ventana, text="ID Paciente").pack(pady=5)
    e_id = tk.Entry(ventana, width=35)
    e_id.pack()

    tk.Label(ventana, text="Genero").pack(pady=5)
    genero_var = tk.StringVar(value="Masculino")
    frame_gen = tk.Frame(ventana)
    frame_gen.pack()
    tk.Radiobutton(frame_gen, text="Masculino", variable=genero_var, value="Masculino").pack(side="left", padx=10)
    tk.Radiobutton(frame_gen, text="Femenino",  variable=genero_var, value="Femenino").pack(side="left", padx=10)
    tk.Radiobutton(frame_gen, text="Otro",      variable=genero_var, value="Otro").pack(side="left", padx=10)

    tk.Label(ventana, text="Edad").pack(pady=5)
    e_edad = tk.Entry(ventana, width=35)
    e_edad.pack()

    # Línea separadora gris
    tk.Frame(ventana, height=1, bg="#cccccc").pack(fill="x", padx=20, pady=10)

    def iniciar():
        """Valida los campos obligatorios y cierra la ventana."""
        if not e_id.get().strip():
            messagebox.showerror("Error", "Introduzca el ID del paciente")
            return
        if not e_edad.get().strip():
            messagebox.showerror("Error", "Introduzca la edad")
            return

        datos["id"]     = e_id.get().strip()
        datos["genero"] = genero_var.get()
        datos["edad"]   = e_edad.get().strip()
        ventana.destroy()

    ttk.Button(ventana, text="Iniciar Evaluacion", command=iniciar).pack(pady=14)
    ventana.mainloop()
    return datos

paciente          = obtener_datos_paciente()
paciente_id       = paciente["id"]
genero            = paciente["genero"]
edad              = paciente["edad"]


# ================= ESTADO DEL JUEGO =================
score               = 0       # nº de taps registrados
game_over           = False   # True cuando se agota el tiempo
finger_tap_detected = False   # estado anterior del tap (para detectar el flanco)
GAME_DURATION       = 20      # duración de la partida (s)
game_start_time     = None    # instante de inicio (None hasta el primer frame)
time_left           = GAME_DURATION

# ================= METRICAS =================
BUFFER_LEN = 200   # nº de puntos visibles en las gráficas en tiempo real

# Buffers circulares (deque con maxlen) para el panel en tiempo real.
dist_history = deque([0.0] * BUFFER_LEN, maxlen=BUFFER_LEN)
vel_history  = deque([0.0] * BUFFER_LEN, maxlen=BUFFER_LEN)

# Historial completo de la partida (para la gráfica PNG final).
dist_history_full = []
vel_history_full  = []
time_history_full = []

prev_dist    = None   # distancia del frame anterior (para derivar la velocidad)
prev_time    = None   # instante del frame anterior
ultimo_tap_t = None   # instante del último tap (para calcular el ITI)

# Una entrada por tap registrado:
tap_distances    = []   # cm: apertura máxima alcanzada ANTES de cada tap
tap_velocities   = []   # cm/s: |velocidad| en el frame del tap
tap_intervals    = []   # s: tiempo entre taps consecutivos (ITI)
tap_timestamps   = []   # s: instante de cada tap desde el inicio de la partida
metricas_finales = {}   # se rellena con calcular_metricas()

temblor_metricas   = None     # dict con freq/amplitud/concluyente por eje (X,Y,Z)
_temblor_procesado = False    # evita detener/calcular el temblor mas de una vez

# Acumulador: apertura máxima alcanzada desde el último tap. Se guarda como
# "amplitud" del tap siguiente y se reinicia a 0 tras cada tap.
_max_dist_entre_taps = 0.0


# ================= FUNCIONES DE DISTANCIA =================
def px_distance(lm, cam_w, cam_h):
    """Distancia en píxeles entre la punta del índice (landmark 8) y la punta
    del pulgar (landmark 4). ``lm`` son los landmarks normalizados de MediaPipe."""
    ix = lm[8].x * cam_w;  iy = lm[8].y * cam_h
    tx = lm[4].x * cam_w;  ty = lm[4].y * cam_h
    return math.hypot(ix - tx, iy - ty)

def px_to_cm(dist_px, lm, cam_w, cam_h):
    """Convierte píxeles a centímetros usando la referencia anatómica
    muñeca(0)-base del dedo medio(9) = HAND_REF_REAL_CM.

    Devuelve 0.0 si la referencia mide < 1 px (mano mal detectada) para evitar
    divisiones por cero."""
    wx = lm[0].x * cam_w;  wy = lm[0].y * cam_h
    mx = lm[9].x * cam_w;  my = lm[9].y * cam_h
    ref_px = math.hypot(mx - wx, my - wy)
    if ref_px < 1:
        return 0.0
    return dist_px * (HAND_REF_REAL_CM / ref_px)

def hand_distance_cm(lm, cam_w, cam_h):
    """Distancia índice-pulgar en centímetros (atajo de las dos anteriores)."""
    return px_to_cm(px_distance(lm, cam_w, cam_h), lm, cam_w, cam_h)


# ================= PAJARO — FISICA =================
# Mecánica: la gravedad tira del pájaro hacia abajo continuamente y cada tap le
# da un impulso hacia arriba. Los valores son bajos a propósito para que el
# movimiento sea lento y legible para pacientes con movilidad reducida.
BIRD_X      = GAME_W // 2                  # posición horizontal fija
BIRD_Y_MIN  =  WINDOW_H // 2               # "techo" (mitad de la pantalla)
BIRD_Y_MAX  = WINDOW_H - 60                # "suelo"
BIRD_Y_INIT = (BIRD_Y_MIN + BIRD_Y_MAX) // 2   # arranca en el centro

bird_y  = float(BIRD_Y_INIT)   # posición vertical actual (px)
bird_vy = 0.0                  # velocidad vertical (px/s, positivo = hacia abajo)

GRAVITY     =  100.0   # px/s² — gravedad constante hacia abajo
TAP_IMPULSE = -80.0   # px/s  — impulso hacia arriba en cada tap
VY_MAX_DOWN =  150.0   # velocidad maxima de caida
VY_MAX_UP   = -100.0   # velocidad maxima de subida

# Estado de la animación del ala (solo estético)
wing_angle = 0.0
wing_dir   = 1


def update_bird(dt):
    """Avanza la física del pájaro ``dt`` segundos (integración de Euler)."""
    global bird_y, bird_vy, wing_angle, wing_dir

    # Fisica: gravedad acumula velocidad hacia abajo
    bird_vy = min(bird_vy + GRAVITY * dt, VY_MAX_DOWN)
    bird_y  = bird_y + bird_vy * dt

    # Limites techo y suelo
    if bird_y < BIRD_Y_MIN:
        bird_y = float(BIRD_Y_MIN); bird_vy = 0.0
    if bird_y > BIRD_Y_MAX:
        bird_y = float(BIRD_Y_MAX); bird_vy = 0.0

    # Animacion ala: bate rapido al subir, lento al caer
    wing_speed = 180.0 if bird_vy < 0 else 60.0
    wing_angle += wing_dir * wing_speed * dt
    if wing_angle >  25.0: wing_angle =  25.0; wing_dir = -1
    if wing_angle < -25.0: wing_angle = -25.0; wing_dir =  1


def draw_bird(frame):
    """Dibuja el pájaro (cola, cuerpo, ala, pico y ojo) y el suelo con
    primitivas de OpenCV. Los colores están en formato BGR."""
    cx = int(BIRD_X)
    cy = int(bird_y)

    # Cola
    tail_pts = np.array([
        [cx - 28, cy],
        [cx - 48, cy - 10],
        [cx - 48, cy + 10]
    ], dtype=np.int32)
    cv2.fillPoly(frame, [tail_pts], (255, 160, 30))

    # Cuerpo
    cv2.ellipse(frame, (cx, cy), (32, 22), 0, 0, 360, (255, 215, 0), -1)
    cv2.ellipse(frame, (cx, cy), (32, 22), 0, 0, 360, (220, 150, 0),  2)

    # Ala (la inclinación depende de wing_angle)
    ala_pts = cv2.ellipse2Poly((cx - 5, cy - 4), (22, 10), int(wing_angle), 0, 360, 5)
    cv2.fillPoly(frame, [ala_pts], (255, 190, 20))
    cv2.polylines(frame, [ala_pts], True, (200, 130, 0), 1)

    # Pico
    pico_pts = np.array([
        [cx + 30, cy - 3],
        [cx + 48, cy],
        [cx + 30, cy + 3]
    ], dtype=np.int32)
    cv2.fillPoly(frame, [pico_pts], (255, 120, 0))

    # Ojo
    cv2.circle(frame, (cx + 14, cy - 7), 7, (255, 255, 255), -1)
    cv2.circle(frame, (cx + 16, cy - 7), 4, (30,  30,  30),  -1)
    cv2.circle(frame, (cx + 17, cy - 9), 1, (255, 255, 255), -1)

    # Suelo visual
    cv2.rectangle(frame, (0, BIRD_Y_MAX), (GAME_W, WINDOW_H), (40, 20, 20), -1)
    cv2.line(frame, (0, BIRD_Y_MAX), (GAME_W, BIRD_Y_MAX), (180, 60, 60), 2)


# ================= REGISTRO DE TAP =================
def register_tap(vel_ms):
    """Registra un tap: guarda sus métricas y da impulso al pájaro.

    Parameters
    ----------
    vel_ms : float
        Velocidad índice-pulgar (cm/s) en el frame en que se detectó el tap.
        Se guarda su valor absoluto como "velocidad de cierre". (Aunque el
        nombre dice ``ms``, la unidad real es cm/s.)
    """
    global ultimo_tap_t, score, _max_dist_entre_taps, bird_vy

    now = time.time()
    # Amplitud = apertura máxima desde el tap anterior.
    tap_distances.append(_max_dist_entre_taps)
    tap_velocities.append(abs(vel_ms))
    if game_start_time:
        tap_timestamps.append(now - game_start_time)

    # ITI (inter-tap interval). Se descartan valores imposibles: < 50 ms
    # (doble detección) o > 5 s (pausa larga, no es un ritmo de tapping).
    if ultimo_tap_t is not None:
        interval = now - ultimo_tap_t
        if 0.05 < interval < 5.0:
            tap_intervals.append(interval)

    ultimo_tap_t         = now
    _max_dist_entre_taps = 0.0     # reinicia el acumulador para el siguiente tap
    score               += 1

    # Impulso hacia arriba (limitado a la velocidad máxima de subida)
    bird_vy = max(bird_vy + TAP_IMPULSE, VY_MAX_UP)


# ================= FIN DE PARTIDA: TEMBLOR =================
def procesar_fin_temblor():
    """Detiene la captura del micro:bit y calcula sus metricas en el
    instante en que termina la partida, para poder mostrarlas ya en la
    pantalla de fin de juego (antes de guardar nada en disco).

    Es idempotente: solo actúa la primera vez que se llama."""
    global temblor_metricas, _temblor_procesado
    if _temblor_procesado:
        return
    _temblor_procesado = True
    temblor.detener()
    temblor_metricas = temblor.calcular_metricas()


# ================= CALCULO DE METRICAS =================
def calcular_metricas():
    """Calcula las métricas clínicas finales a partir de las listas por tap.

    Devuelve un diccionario con:
        amp_media / vel_media   : medias de amplitud (cm) y velocidad (cm/s)
        pendiente_amp / _vel    : pendiente de la recta de regresión lineal
                                  (unidad por nº de tap)
        decrece_amp / _vel      : True si la pendiente es negativa
                                  (decremento a lo largo de la tarea)
        variabilidad_pct        : coeficiente de variación de la amplitud (%)
                                  = desviación típica / media * 100
    """
    metricas = {}
    n  = len(tap_distances)
    nv = len(tap_velocities)

    metricas["amp_media"] = float(np.mean(tap_distances))  if n  >= 1 else 0.0
    metricas["vel_media"] = float(np.mean(tap_velocities)) if nv >= 1 else 0.0

    # Pendiente de la amplitud: ajuste lineal amplitud = a * nº_tap + b.
    # Se necesitan al menos 2 taps para ajustar una recta.
    if n >= 2:
        x = np.arange(n, dtype=float)
        pend_a, _ = np.polyfit(x, tap_distances, 1)
        metricas["pendiente_amp"] = float(pend_a)
        metricas["decrece_amp"]   = pend_a < 0
    else:
        metricas["pendiente_amp"] = 0.0
        metricas["decrece_amp"]   = False

    # Pendiente de la velocidad (mismo procedimiento)
    if nv >= 2:
        x = np.arange(nv, dtype=float)
        pend_v, _ = np.polyfit(x, tap_velocities, 1)
        metricas["pendiente_vel"] = float(pend_v)
        metricas["decrece_vel"]   = pend_v < 0
    else:
        metricas["pendiente_vel"] = 0.0
        metricas["decrece_vel"]   = False

    # Variabilidad (CV%): ddof=1 -> desviación típica muestral.
    if n >= 2 and metricas["amp_media"] > 0:
        metricas["variabilidad_pct"] = (float(np.std(tap_distances, ddof=1)) / metricas["amp_media"]) * 100.0
    else:
        metricas["variabilidad_pct"] = 0.0

    return metricas


# ================= GUARDAR GRAFICAS PNG =================
def guardar_graficas():
    """Genera un PNG 2x2 con las gráficas de la sesión:

        [0,0] amplitud por tap + línea de tendencia + media
        [0,1] velocidad por tap + línea de tendencia + media
        [1,0] distancia índice-pulgar continua (con los taps marcados en rojo)
        [1,1] velocidad continua (rojo = cierre, verde = apertura)
    """
    global metricas_finales
    metricas_finales = calcular_metricas()
    m = metricas_finales

    nombre_png = f"grafica_fingertap_{paciente_id}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.png"
    fig, axes  = plt.subplots(2, 2, figsize=(14, 8))
    fig.suptitle(f"Sesion Fingertap  |  ID: {paciente_id} | Genero: {genero} | Edad: {edad}",
                 fontsize=13, fontweight='bold')

    # Amplitud por tap
    ax1 = axes[0, 0]
    if tap_distances:
        x_t = np.arange(len(tap_distances))
        ax1.plot(x_t, tap_distances, color='#4db8ff', linewidth=1.2,
                 marker='o', markersize=4, label='Amplitud por tap')
        if len(tap_distances) >= 2:
            coef = np.polyfit(x_t, tap_distances, 1)
            ax1.plot(x_t, np.polyval(coef, x_t),
                     color='#ff4444', linewidth=1.5, linestyle='--', label='Tendencia')
        ax1.axhline(m["amp_media"], color='yellow', linewidth=0.9, linestyle=':',
                    alpha=0.8, label=f'Media={m["amp_media"]:.2f}cm')
    ax1.set_title("Amplitud de movimiento por tap")
    ax1.set_ylabel("Distancia (cm)"); ax1.set_xlabel("Tap #")
    ax1.legend(fontsize=7); ax1.grid(True, alpha=0.3)

    # Velocidad por tap
    ax2 = axes[0, 1]
    if tap_velocities:
        x_t = np.arange(len(tap_velocities))
        ax2.plot(x_t, tap_velocities, color='#ffa040', linewidth=1.2,
                 marker='s', markersize=4, label='Velocidad por tap')
        if len(tap_velocities) >= 2:
            coef = np.polyfit(x_t, tap_velocities, 1)
            ax2.plot(x_t, np.polyval(coef, x_t),
                     color='#ff4444', linewidth=1.5, linestyle='--', label='Tendencia')
        ax2.axhline(m["vel_media"], color='yellow', linewidth=0.9, linestyle=':',
                    alpha=0.8, label=f'Media={m["vel_media"]:.2f}cm/s')
    ax2.set_title("Velocidad de cierre por tap")
    ax2.set_ylabel("Velocidad (cm/s)"); ax2.set_xlabel("Tap #")
    ax2.legend(fontsize=7); ax2.grid(True, alpha=0.3)

    # Distancia continua (n3 = longitud común por si las listas difieren en 1)
    ax3 = axes[1, 0]
    if time_history_full and dist_history_full:
        n3 = min(len(time_history_full), len(dist_history_full))
        ax3.plot(time_history_full[:n3], dist_history_full[:n3],
                 color='#4db8ff', linewidth=0.8, label='Distancia (cm)')
        if tap_timestamps and tap_distances:
            ax3.scatter(tap_timestamps, tap_distances,
                        color='red', s=20, zorder=5, label='Tap')
    ax3.set_title("Distancia indice-pulgar (continua)")
    ax3.set_ylabel("Distancia (cm)"); ax3.set_xlabel("Tiempo (s)")
    ax3.legend(fontsize=7); ax3.grid(True, alpha=0.3)
    ax3.set_xlim(left=0); ax3.set_ylim(bottom=0)

    # Velocidad continua (negativa = cierre, positiva = apertura)
    ax4 = axes[1, 1]
    if time_history_full and vel_history_full:
        n4 = min(len(time_history_full), len(vel_history_full))
        t  = time_history_full[:n4]
        v  = vel_history_full[:n4]
        ax4.plot(t, v, color='#ffa040', linewidth=0.8, label='Velocidad (cm/s)')
        ax4.axhline(0, color='white', linewidth=0.6, linestyle='--', alpha=0.5)
        ax4.fill_between(t, v, 0, where=[vv < 0 for vv in v],
                         color='#ff6060', alpha=0.25, label='Cierre')
        ax4.fill_between(t, v, 0, where=[vv > 0 for vv in v],
                         color='#60ff80', alpha=0.25, label='Apertura')
    ax4.set_title("Velocidad cierre/apertura (continua)")
    ax4.set_ylabel("Velocidad (cm/s)"); ax4.set_xlabel("Tiempo (s)")
    ax4.legend(fontsize=7); ax4.grid(True, alpha=0.3)
    ax4.set_xlim(left=0)

    plt.tight_layout()
    plt.savefig(nombre_png, dpi=150, bbox_inches='tight')
    plt.close(fig)
    print("Grafica guardada:", nombre_png)


# ================= GUARDAR EXCEL =================
def guardar_resultados():
    """Añade una fila con los resultados de la sesión a ``sesiones_fingertap.xlsx``.

    Si el archivo no existe se crea con su fila de cabecera; si existe, se
    añade al final (así se acumulan todas las sesiones en un solo Excel).
    Las columnas de temblor quedan vacías si no hubo datos del sensor.
    """
    global metricas_finales
    if not metricas_finales:
        metricas_finales = calcular_metricas()
    m      = metricas_finales
    fecha  = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    archivo = "sesiones_fingertap.xlsx"

    try:
        wb = load_workbook(archivo)
        ws = wb.active
    except:  # noqa: E722 - el archivo no existe aún: se crea uno nuevo
        wb = Workbook()
        ws = wb.active
        ws.title = "sesiones_fingertap"
        ws.append(["Fecha", "ID", "Genero", "Edad", "Taps",
                   "AmpMedia_cm", "VelMedia_cms",
                   "PendienteAmp", "DecreceAmp",
                   "PendienteVel", "DecreceVel",
                   "Variabilidad_pct",
                   "Temblor_FreqX_Hz", "Temblor_AmpX", "Temblor_ConcluyenteX",
                   "Temblor_FreqY_Hz", "Temblor_AmpY", "Temblor_ConcluyenteY",
                   "Temblor_FreqZ_Hz", "Temblor_AmpZ", "Temblor_ConcluyenteZ"])

    fila = [
        fecha, paciente_id, genero, edad, score,
        round(m["amp_media"],        2),
        round(m["vel_media"],        2),
        round(m["pendiente_amp"],    6),
        "SI" if m["decrece_amp"] else "NO",
        round(m["pendiente_vel"],    6),
        "SI" if m["decrece_vel"] else "NO",
        round(m["variabilidad_pct"], 2)
    ]

    # Tres columnas por eje (frecuencia, amplitud, concluyente); None si no hay dato.
    for eje in ("X", "Y", "Z"):
        if temblor_metricas and temblor_metricas.get(eje, {}).get("freq_hz") is not None:
            d = temblor_metricas[eje]
            fila.append(round(d["freq_hz"], 2))
            fila.append(round(d["amplitud"], 1))
            fila.append("SI" if d["concluyente"] else "NO")
        else:
            fila.extend([None, None, None])

    ws.append(fila)
    wb.save(archivo)
    print("Excel guardado:", archivo)


# ================= HELPERS DE DIBUJO =================
def put_text(frame, text, pos, scale=0.8, color=(255, 255, 255), thickness=2):
    """Escribe texto en el frame. OpenCV no soporta tildes ni 'ñ' con las
    fuentes Hershey, así que se convierten a '?' para evitar caracteres rotos."""
    safe = text.encode('ascii', 'replace').decode('ascii')
    cv2.putText(frame, safe, pos, cv2.FONT_HERSHEY_SIMPLEX,
                scale, color, thickness, cv2.LINE_AA)


def draw_temblor_block(frame, x, y, scale=0.62):
    """Dibuja frecuencia y amplitud por eje (X, Y, Z) del temblor.

    Verde = pico concluyente; ocre = no concluyente (probable ruido);
    gris = sin datos. Devuelve la coordenada y tras el bloque dibujado."""
    put_text(frame, "-- Temblor (acelerometro) --", (x, y),
             scale=0.55, color=(160, 160, 180), thickness=1); y += 26

    if not temblor.conectado:
        put_text(frame, "Sensor no conectado", (x, y),
                 scale=scale, color=(150, 150, 150), thickness=1); y += 28
        return y

    if temblor_metricas is None:
        put_text(frame, "Datos insuficientes", (x, y),
                 scale=scale, color=(150, 150, 150), thickness=1); y += 28
        return y

    for eje in ("X", "Y", "Z"):
        d = temblor_metricas[eje]
        if d["freq_hz"] is None:
            texto = f"Eje {eje}: sin datos"
            color = (120, 120, 120)
        else:
            texto = f"Eje {eje}: {d['freq_hz']:.2f} Hz   amp {d['amplitud']:.0f}"
            if not d["concluyente"]:
                texto += "  (no concluyente / ruido)"
            color = (80, 220, 80) if d["concluyente"] else (180, 160, 90)
        put_text(frame, texto, (x, y), scale=scale, color=color, thickness=1)
        y += 26
    return y


def draw_graph(frame, data, graph_x, graph_y, graph_w, graph_h,
               color, label, unit, y_min=None, y_max=None, zero_line=False):
    """Dibuja una gráfica de línea sencilla (estilo osciloscopio) con OpenCV.

    Parameters
    ----------
    data : iterable de float
        Valores a representar (normalmente un deque circular).
    graph_x, graph_y, graph_w, graph_h : int
        Posición y tamaño del recuadro de la gráfica en píxeles.
    y_min, y_max : float | None
        Límites del eje Y; si son None se autoescalan con los datos.
    zero_line : bool
        Si True y el rango cruza el 0, dibuja la línea del cero.
    """
    # Fondo
    cv2.rectangle(frame, (graph_x, graph_y),
                  (graph_x + graph_w, graph_y + graph_h), (40, 40, 55), -1)
    vals = list(data)
    if not vals:
        return
    lo  = y_min if y_min is not None else min(vals)
    hi  = y_max if y_max is not None else max(vals)
    rng = hi - lo if hi != lo else 1.0   # evita dividir por cero con señal plana

    # Línea del cero
    if zero_line and lo < 0 < hi:
        yz = graph_y + graph_h - int((0 - lo) / rng * graph_h)
        cv2.line(frame, (graph_x, yz), (graph_x + graph_w, yz), (80, 80, 100), 1)

    # Rejilla horizontal al 25 %, 50 % y 75 % con su valor
    for frac in [0.25, 0.5, 0.75]:
        yg = graph_y + int((1 - frac) * graph_h)
        cv2.line(frame, (graph_x, yg), (graph_x + graph_w, yg), (55, 55, 70), 1)
        put_text(frame, f"{lo + frac * rng:.3f}",
                 (graph_x + graph_w - 45, yg - 2),
                 scale=0.3, color=(90, 90, 110), thickness=1)

    # Conversión de cada valor a coordenadas de píxel (el eje Y está invertido
    # en las imágenes: y=0 es la parte superior).
    step = graph_w / max(len(vals) - 1, 1)
    pts  = []
    for i, v in enumerate(vals):
        gx = int(graph_x + i * step)
        gy = int(graph_y + graph_h - ((v - lo) / rng) * graph_h)
        gy = max(graph_y, min(graph_y + graph_h, gy))   # recorta al recuadro
        pts.append((gx, gy))
    for i in range(1, len(pts)):
        cv2.line(frame, pts[i - 1], pts[i], color, 1)

    # Borde y valor actual
    cv2.rectangle(frame, (graph_x, graph_y),
                  (graph_x + graph_w, graph_y + graph_h), (90, 90, 120), 1)
    put_text(frame, f"{label}: {vals[-1]:.4f} {unit}",
             (graph_x, graph_y - 5), scale=0.48, color=color, thickness=1)


# ================= PANELES =================
def draw_panel(frame):
    """Panel lateral DURANTE la partida: ITI medio, velocidad media y las
    gráficas en tiempo real de distancia y velocidad."""
    px = GAME_W
    cv2.rectangle(frame, (px, 0), (WINDOW_W, WINDOW_H), (25, 25, 35), -1)
    cv2.line(frame, (px, 0), (px, WINDOW_H), (70, 70, 100), 2)

    put_text(frame, "FINGERTAP MONITOR", (px + 15, 32), scale=0.72, color=(180, 220, 255))

    y = 58
    if tap_intervals:
        avg_int = sum(tap_intervals) / len(tap_intervals)
        put_text(frame, f"ITI medio: {avg_int:.3f} s", (px+12, y),
                 scale=0.58, color=(200, 255, 200), thickness=1)
    else:
        put_text(frame, "ITI medio: --", (px+12, y),
                 scale=0.58, color=(140, 140, 140), thickness=1)
    y += 22

    if tap_velocities:
        avg_vel = sum(tap_velocities) / len(tap_velocities)
        put_text(frame, f"Vel. cierre media: {avg_vel:.2f} cm/s", (px+12, y),
                 scale=0.58, color=(255, 210, 120), thickness=1)
    else:
        put_text(frame, "Vel. cierre media: --", (px+12, y),
                 scale=0.58, color=(140, 140, 140), thickness=1)
    y += 22

    if tap_distances:
        put_text(frame, "Distancia tap registrada", (px+12, y),
                 scale=0.58, color=(130, 200, 255), thickness=1)
    else:
        put_text(frame, "Distancia tap: --", (px+12, y),
                 scale=0.58, color=(140, 140, 140), thickness=1)
    y += 22

    cv2.line(frame, (px+10, y), (WINDOW_W-10, y), (60, 60, 80), 1)
    y += 12

    gw = PANEL_W - 28; gh = 110; gx = px + 14

    put_text(frame, "Distancia indice-pulgar (cm):", (gx, y),
             scale=0.48, color=(160, 200, 255), thickness=1)
    y += 14
    draw_graph(frame, dist_history, gx, y, gw, gh,
               color=(80, 180, 255), label="Dist", unit="cm", y_min=0)
    y += gh + 22

    put_text(frame, "Velocidad cierre/apertura (cm/s):", (gx, y),
             scale=0.48, color=(255, 200, 100), thickness=1)
    y += 14
    draw_graph(frame, vel_history, gx, y, gw, gh,
               color=(255, 160, 40), label="Vel", unit="cm/s", zero_line=True)
    y += gh + 15

    cv2.line(frame, (px+10, y), (WINDOW_W-10, y), (60, 60, 80), 1)
    y += 8
    put_text(frame, "Vel > 0: apertura  |  Vel < 0: cierre", (px+10, y),
             scale=0.38, color=(120, 120, 140), thickness=1)

    # FPS medio de los últimos 30 frames (fps_buffer se define más abajo,
    # antes del bucle principal)
    fps_val = sum(fps_buffer) / len(fps_buffer) if fps_buffer else 0
    put_text(frame, f"FPS: {fps_val:.0f}", (px+12, WINDOW_H - 15),
             scale=0.48, color=(80, 80, 110), thickness=1)


def draw_panel_results(frame):
    """Panel lateral AL TERMINAR: resumen de métricas (amplitud, velocidad,
    tendencias, variabilidad con barra y temblor)."""
    global metricas_finales
    if not metricas_finales:
        metricas_finales = calcular_metricas()
    m  = metricas_finales
    px = GAME_W

    cv2.rectangle(frame, (px, 0), (WINDOW_W, WINDOW_H), (25, 25, 35), -1)
    cv2.line(frame, (px, 0), (px, WINDOW_H), (70, 70, 100), 2)

    put_text(frame, "RESULTADOS", (px + 30, 38), scale=0.9, color=(34, 197, 94), thickness=2)

    y = 65
    cv2.line(frame, (px+10, y), (WINDOW_W-10, y), (60, 60, 80), 1); y += 16
    put_text(frame, f"Taps totales: {score}", (px+15, y), scale=0.72, color=(180, 220, 255))
    y += 32

    cv2.line(frame, (px+10, y), (WINDOW_W-10, y), (60, 60, 80), 1); y += 14
    put_text(frame, "-- Amplitud de movimiento --", (px+12, y),
             scale=0.55, color=(160, 160, 180), thickness=1); y += 26
    put_text(frame, f"Amplitud media: {m['amp_media']:.2f} cm",
             (px+12, y), scale=0.68, color=(130, 200, 255)); y += 32

    cv2.line(frame, (px+10, y), (WINDOW_W-10, y), (60, 60, 80), 1); y += 14
    put_text(frame, "-- Velocidad de cierre --", (px+12, y),
             scale=0.55, color=(160, 160, 180), thickness=1); y += 26
    put_text(frame, f"Velocidad media: {m['vel_media']:.2f} cm/s",
             (px+12, y), scale=0.68, color=(255, 210, 100)); y += 32

    cv2.line(frame, (px+10, y), (WINDOW_W-10, y), (60, 60, 80), 1); y += 14
    put_text(frame, "-- Tendencia (pendiente) --", (px+12, y),
             scale=0.55, color=(160, 160, 180), thickness=1); y += 26

    # Verde = sin decremento; rojo (BGR) = la pendiente es negativa.
    col_a = (80, 220, 80) if not m["decrece_amp"] else (80, 80, 255)
    put_text(frame, f"Amp: {m['pendiente_amp']:+.5f}  [{'SI DECRECE' if m['decrece_amp'] else 'NO decrece'}]",
             (px+12, y), scale=0.60, color=col_a); y += 28
    col_v = (80, 220, 80) if not m["decrece_vel"] else (80, 80, 255)
    put_text(frame, f"Vel: {m['pendiente_vel']:+.5f}  [{'SI DECRECE' if m['decrece_vel'] else 'NO decrece'}]",
             (px+12, y), scale=0.60, color=col_v); y += 34

    cv2.line(frame, (px+10, y), (WINDOW_W-10, y), (60, 60, 80), 1); y += 14
    put_text(frame, "-- Variabilidad --", (px+12, y),
             scale=0.55, color=(160, 160, 180), thickness=1); y += 26
    # Umbral visual del 20 %: por debajo, verde; por encima, naranja.
    vc = (200, 255, 180) if m["variabilidad_pct"] < 20 else (255, 160, 80)
    put_text(frame, f"Variabilidad: {m['variabilidad_pct']:.1f} %",
             (px+12, y), scale=0.72, color=vc); y += 28

    # Barra de variabilidad: la barra completa equivale al 50 %.
    bx = px + 12; bw = PANEL_W - 28; bh = 10
    pct_c = min(m["variabilidad_pct"] / 50.0, 1.0)
    cv2.rectangle(frame, (bx, y), (bx + bw, y + bh), (60, 60, 80), -1)
    fc = (80, 200, 80) if m["variabilidad_pct"] < 20 else (255, 160, 50)
    cv2.rectangle(frame, (bx, y), (bx + int(bw * pct_c), y + bh), fc, -1)
    cv2.rectangle(frame, (bx, y), (bx + bw, y + bh), (120, 120, 150), 1)
    y += bh + 16

    cv2.line(frame, (px+10, y), (WINDOW_W-10, y), (60, 60, 80), 1); y += 14
    draw_temblor_block(frame, px+12, y, scale=0.5)


def draw_top_bar(frame):
    """Barra superior: título, datos del paciente, tiempo restante, taps y
    rango de ITI (mínimo y máximo)."""
    cv2.rectangle(frame, (0, 0), (GAME_W, 65), (21, 94, 88), -1)
    put_text(frame, "Juego de Rehabilitacion", (20, 42), scale=1.0)
    put_text(frame, f"ID: {paciente_id} | Genero: {genero} | Edad: {edad}",
             (20, 62), scale=0.45)
    put_text(frame, f"Tiempo: {max(time_left, 0)}s", (GAME_W - 200, 42), scale=0.9)
    score_str = f"Taps: {score}"
    if tap_intervals:
        mn = min(tap_intervals); mx = max(tap_intervals)
        score_str += f"   ITI Min: {mn:.3f}s  Max: {mx:.3f}s"
    put_text(frame, score_str, (20, 95), scale=0.65, color=(20, 150, 130))


def draw_game_over(frame):
    """Pantalla de fin de partida: oscurece el juego y muestra el resumen
    de métricas sobre él."""
    global metricas_finales
    if not metricas_finales:
        metricas_finales = calcular_metricas()
    m = metricas_finales

    # Oscurecer el fondo con una capa negra semitransparente (60 %).
    overlay = frame.copy()
    cv2.rectangle(overlay, (0, 0), (GAME_W, WINDOW_H), (0, 0, 0), -1)
    cv2.addWeighted(overlay, 0.6, frame, 0.4, 0, frame)
    cx = GAME_W // 2

    put_text(frame, "Tiempo terminado!", (cx-200, WINDOW_H//2-150),
             scale=1.2, color=(34, 197, 94), thickness=3)
    put_text(frame, f"Taps: {score}", (cx-60, WINDOW_H//2-112), scale=0.95)

    y = WINDOW_H // 2 - 78
    put_text(frame, f"Amplitud media:   {m['amp_media']:.2f} cm",
             (cx-235, y), scale=0.75, color=(130, 200, 255)); y += 34
    put_text(frame, f"Velocidad media:  {m['vel_media']:.2f} cm/s",
             (cx-235, y), scale=0.75, color=(255, 210, 120)); y += 34
    col_a = (80, 220, 80) if not m["decrece_amp"] else (80, 100, 255)
    put_text(frame, f"Tend. amplitud:   {m['pendiente_amp']:+.5f}  ({'DECRECE' if m['decrece_amp'] else 'estable'})",
             (cx-235, y), scale=0.72, color=col_a); y += 32
    col_v = (80, 220, 80) if not m["decrece_vel"] else (80, 100, 255)
    put_text(frame, f"Tend. velocidad:  {m['pendiente_vel']:+.5f}  ({'DECRECE' if m['decrece_vel'] else 'estable'})",
             (cx-235, y), scale=0.72, color=col_v); y += 34
    vc = (200, 255, 160) if m["variabilidad_pct"] < 20 else (255, 160, 80)
    put_text(frame, f"Variabilidad:     {m['variabilidad_pct']:.1f} %",
             (cx-235, y), scale=0.75, color=vc); y += 40

    y = draw_temblor_block(frame, cx-235, y, scale=0.65)

    put_text(frame, "Pulsa Q para salir", (cx-130, y),
             scale=0.72, color=(150, 150, 150), thickness=1)


# ================= FPS =================
fps_buffer   = deque(maxlen=30)   # FPS instantáneos de los últimos 30 frames
last_frame_t = time.time()

# ================= LOOP PRINCIPAL =================
print("Iniciando... Pulsa Q para salir.")

while True:
    now_t = time.time()

    # --- 1) Captura del frame -------------------------------------------------
    ret, cam_frame = cap.read()
    if not ret:
        print("ERROR: No se puede leer la camara.")
        break

    # dt = tiempo desde el frame anterior (lo usa la física del pájaro)
    dt = now_t - last_frame_t
    fps_buffer.append(1.0 / max(dt, 1e-6))
    last_frame_t = now_t

    # Efecto espejo: así el movimiento en pantalla coincide con el de la mano.
    cam_frame = cv2.flip(cam_frame, 1)
    cam_h, cam_w = cam_frame.shape[:2]

    # MediaPipe espera RGB; OpenCV trabaja en BGR.
    rgb     = cv2.cvtColor(cam_frame, cv2.COLOR_BGR2RGB)
    results = hands.process(rgb)

    vel = 0.0   # velocidad índice-pulgar de este frame (cm/s)

    # --- 2) Medida de la mano -------------------------------------------------
    if results.multi_hand_landmarks:
        # Dibuja el esqueleto de la mano sobre la imagen de cámara.
        for hl in results.multi_hand_landmarks:
            mp_draw.draw_landmarks(
                cam_frame, hl,
                mp_hands.HAND_CONNECTIONS,
                mp_draw_styles.get_default_hand_landmarks_style(),
                mp_draw_styles.get_default_hand_connections_style()
            )

        lm     = results.multi_hand_landmarks[0].landmark
        dist_cm = hand_distance_cm(lm, cam_w, cam_h)
        dist_history.append(dist_cm)

        # Solo se acumulan datos si la partida está en curso.
        if game_start_time and not game_over:
            dist_history_full.append(dist_cm)
            time_history_full.append(now_t - game_start_time)
            # Apertura máxima desde el último tap (= amplitud del próximo tap)
            if dist_cm > _max_dist_entre_taps:
                _max_dist_entre_taps = dist_cm

        # Velocidad = derivada numérica de la distancia entre dos frames.
        # Negativa = los dedos se acercan (cierre); positiva = se separan.
        if prev_dist is not None and prev_time is not None:
            dt_hand = now_t - prev_time
            if dt_hand > 0:
                vel = (dist_cm - prev_dist) / dt_hand
            vel_history.append(vel)
            if game_start_time and not game_over:
                vel_history_full.append(vel)
        else:
            vel_history.append(0.0)
            if game_start_time and not game_over:
                vel_history_full.append(0.0)

        prev_dist = dist_cm
        prev_time = now_t

        # --- 3) Detección de tap ----------------------------------------------
        # Un tap es el FLANCO de subida del estado "dedos juntos" (< 1 cm):
        # solo se cuenta al pasar de separados a juntos, no mientras se
        # mantienen juntos.
        is_tap = dist_cm < 1.0 #DISTANCIA (umbral de tap en cm)
        if is_tap and not finger_tap_detected and not game_over:
            register_tap(vel)
        finger_tap_detected = is_tap

    else:
        # Mano no detectada: se rellenan ceros para mantener alineadas las series.
        dist_history.append(0.0)
        vel_history.append(0.0)
        if game_start_time and not game_over:
            dist_history_full.append(0.0)
            vel_history_full.append(0.0)
            time_history_full.append(now_t - game_start_time)
        finger_tap_detected = False

    # --- 4) Física del pájaro -------------------------------------------------
    if game_start_time and not game_over:
        update_bird(dt)

    # --- 5) Dibujo ------------------------------------------------------------
    # Lienzo base (BGR) sobre el que se dibuja todo cada frame.
    frame = np.full((WINDOW_H, WINDOW_W, 3), (30, 40, 60), dtype=np.uint8)

    # Inicio de la partida: se arranca el sensor y DESPUÉS el cronómetro.
    if game_start_time is None:
        temblor.iniciar()               # bloquea ~2 s (arranque del micro:bit)
        game_start_time = time.time()   # el cronometro empieza DESPUES de la espera
        last_frame_t = time.time()      # evita que el pajaro "caiga" de golpe por el dt de 2 s
    if not game_over:
        elapsed   = time.time() - game_start_time
        time_left = max(0, GAME_DURATION - int(elapsed))
        if time_left <= 0:
            game_over = True

    # En cuanto acaba la partida se cierra la captura de temblor y se calculan
    # las métricas (para mostrarlas ya en la pantalla final).
    if game_over:
        procesar_fin_temblor()

    if not game_over:
        put_text(frame, "Haz FINGERTAP: une indice y pulgar",
                 (GAME_W//2 - 265, 115), scale=0.75, color=(100, 140, 160))
        draw_bird(frame)

    draw_top_bar(frame)

    if game_over:
        draw_panel_results(frame)
        draw_game_over(frame)
    else:
        draw_panel(frame)

    # Cámara en vivo (miniatura en la esquina inferior izquierda)
    cam_disp_w  = 240
    cam_disp_h  = int(cam_disp_w * cam_h / cam_w)   # mantiene la proporción
    cam_resized = cv2.resize(cam_frame, (cam_disp_w, cam_disp_h))
    y1 = WINDOW_H - cam_disp_h - 10; y2 = WINDOW_H - 10
    x1 = 10;                          x2 = 10 + cam_disp_w
    frame[y1:y2, x1:x2] = cam_resized
    cv2.rectangle(frame, (x1-2, y1-2), (x2+2, y2+2), (21, 94, 88), 2)
    put_text(frame, "Camara en vivo", (x1, y1-6),
             scale=0.5, color=(21, 94, 88), thickness=1)

    cv2.imshow(WIN_NAME, frame)
    # waitKey(1) procesa los eventos de la ventana; 'q' termina el programa.
    if cv2.waitKey(1) & 0xFF == ord('q'):
        break

# ================= CIERRE =================
# Se cierra el sensor por si la ventana se cerró antes de ver la pantalla final.
procesar_fin_temblor()   # por si se cerro la ventana antes de ver la pantalla de fin

# Guardado de resultados en disco (solo ocurre al salir del bucle).
if GUARDAR_CSV_TEMBLOR:
    temblor.guardar_csv(paciente_id)
temblor.guardar_grafica(paciente_id)

guardar_resultados()
guardar_graficas()

# Liberar recursos
cap.release()
cv2.destroyAllWindows()
hands.close()

# Resumen final en consola
print("\n--- RESULTADOS FINALES ---")
print(f"ID: {paciente_id} | Genero: {genero} | Edad: {edad}")
print(f"Taps: {score}")
if tap_intervals:
    avg = sum(tap_intervals) / len(tap_intervals)
    print(f"ITI  min={min(tap_intervals):.3f}s  max={max(tap_intervals):.3f}s  med={avg:.3f}s")
if tap_velocities:
    avg_v = sum(tap_velocities) / len(tap_velocities)
    print(f"Vel. cierre  media={avg_v:.2f} cm/s  max={max(tap_velocities):.2f} cm/s")
if tap_distances:
    print(f"Dist. en tap  min={min(tap_distances):.2f}cm  max={max(tap_distances):.2f}cm")
if temblor_metricas:
    for eje in ("X", "Y", "Z"):
        d = temblor_metricas[eje]
        if d["freq_hz"] is not None:
            nota = "" if d["concluyente"] else "  (no concluyente / ruido)"
            print(f"Temblor eje {eje}: {d['freq_hz']:.2f} Hz  amp={d['amplitud']:.0f}{nota}")
